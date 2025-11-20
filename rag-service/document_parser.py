"""
Document parser for various file formats (PDF, DOCX, TXT)
"""
import io
from typing import Dict, Any
from docx import Document
import pypdf
import pdfplumber
import requests
from bs4 import BeautifulSoup
import json


class DocumentParser:
    """Parse documents from various file formats"""
    
    @staticmethod
    def parse_pdf(file_content: bytes) -> str:
        """
        Parse PDF file and extract text with enhanced table handling
        
        Args:
            file_content: PDF file content as bytes
            
        Returns:
            Extracted text from PDF with formatted tables
        """
        text_parts = []
        
        try:
            # Try with pdfplumber first (better for complex PDFs and tables)
            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    # Extract regular text
                    page_text = page.extract_text()
                    
                    # Extract tables separately for better formatting
                    tables = page.extract_tables()
                    
                    if tables:
                        # If page has tables, format them nicely
                        page_parts = []
                        if page_text:
                            page_parts.append(page_text)
                        
                        for table_idx, table in enumerate(tables, 1):
                            # Format table as markdown-style
                            table_text = f"\n[Table {table_idx} on Page {page_num}]\n"
                            for row in table:
                                if row:
                                    # Clean and join cells
                                    cleaned_row = [str(cell).strip() if cell else "" for cell in row]
                                    table_text += " | ".join(cleaned_row) + "\n"
                            page_parts.append(table_text)
                        
                        text_parts.append("\n".join(page_parts))
                    elif page_text:
                        # No tables, just add text
                        text_parts.append(page_text)
                        
        except Exception as e:
            print(f"pdfplumber failed, trying pypdf: {e}")
            # Fallback to pypdf
            try:
                pdf_reader = pypdf.PdfReader(io.BytesIO(file_content))
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            except Exception as e2:
                raise Exception(f"Failed to parse PDF: {e2}")
        
        return "\n\n".join(text_parts)
    
    @staticmethod
    def parse_docx(file_content: bytes) -> str:
        """
        Parse DOCX file and extract text
        
        Args:
            file_content: DOCX file content as bytes
            
        Returns:
            Extracted text from DOCX
        """
        try:
            doc = Document(io.BytesIO(file_content))
            text_parts = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text for cell in row.cells])
                    if row_text.strip():
                        text_parts.append(row_text)
            
            return "\n\n".join(text_parts)
        except Exception as e:
            raise Exception(f"Failed to parse DOCX: {e}")
    
    @staticmethod
    def parse_txt(file_content: bytes) -> str:
        """
        Parse TXT file
        
        Args:
            file_content: TXT file content as bytes
            
        Returns:
            Text content
        """
        try:
            # Try UTF-8 first
            return file_content.decode('utf-8')
        except UnicodeDecodeError:
            # Fallback to latin-1
            try:
                return file_content.decode('latin-1')
            except Exception as e:
                raise Exception(f"Failed to parse TXT: {e}")
    
    @staticmethod
    def parse_json(file_content: bytes) -> str:
        """
        Parse JSON file and convert to formatted text
        
        Args:
            file_content: JSON file content as bytes
            
        Returns:
            Formatted text from JSON
        """
        try:
            # Decode bytes to string
            json_str = file_content.decode('utf-8')
            data = json.loads(json_str)
            
            # Handle different JSON structures
            if isinstance(data, dict):
                # Single JSON object
                return DocumentParser._format_json_object(data)
            elif isinstance(data, list):
                # Array of JSON objects
                formatted_parts = []
                for idx, item in enumerate(data, 1):
                    if isinstance(item, dict):
                        formatted_parts.append(f"--- Item {idx} ---")
                        formatted_parts.append(DocumentParser._format_json_object(item))
                    else:
                        formatted_parts.append(f"Item {idx}: {str(item)}")
                return "\n\n".join(formatted_parts)
            else:
                # Fallback: convert to string
                return json.dumps(data, indent=2, ensure_ascii=False)
        except json.JSONDecodeError as e:
            raise Exception(f"Invalid JSON format: {e}")
        except Exception as e:
            raise Exception(f"Failed to parse JSON: {e}")
    
    @staticmethod
    def _format_json_object(obj: dict) -> str:
        """
        Format a JSON object into readable text
        
        Args:
            obj: Dictionary object
            
        Returns:
            Formatted text
        """
        parts = []
        
        # Common field names to prioritize
        priority_fields = ['id', 'title', 'name', 'heading']
        content_fields = ['content', 'text', 'body', 'description']
        
        # Add title/heading if exists
        for field in priority_fields:
            if field in obj:
                parts.append(f"# {obj[field]}")
                break
        
        # Add content if exists
        for field in content_fields:
            if field in obj:
                parts.append(obj[field])
                break
        
        # Add remaining fields
        for key, value in obj.items():
            if key not in priority_fields + content_fields:
                if isinstance(value, (dict, list)):
                    parts.append(f"**{key}:**\n{json.dumps(value, indent=2, ensure_ascii=False)}")
                else:
                    parts.append(f"**{key}:** {value}")
        
        return "\n\n".join(parts)
    
    @staticmethod
    def parse_file(file_content: bytes, filename: str) -> Dict[str, Any]:
        """
        Parse file based on extension
        
        Args:
            file_content: File content as bytes
            filename: Original filename
            
        Returns:
            Dictionary with parsed content and metadata
        """
        extension = filename.lower().split('.')[-1]
        
        try:
            if extension == 'pdf':
                content = DocumentParser.parse_pdf(file_content)
            elif extension in ['docx', 'doc']:
                content = DocumentParser.parse_docx(file_content)
            elif extension in ['txt', 'md', 'csv']:
                content = DocumentParser.parse_txt(file_content)
            elif extension == 'json':
                content = DocumentParser.parse_json(file_content)
            else:
                raise Exception(f"Unsupported file format: {extension}")
            
            return {
                "content": content,
                "filename": filename,
                "extension": extension,
                "size": len(file_content),
                "char_count": len(content)
            }
        except Exception as e:
            raise Exception(f"Error parsing {filename}: {str(e)}")

    @staticmethod
    def parse_url(url: str) -> Dict[str, Any]:
        """
        Parse a URL and extract its main content
        
        Args:
            url: The URL to scrape
            
        Returns:
            Dictionary with parsed content and metadata
        """
        try:
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'
            }
            response = requests.get(url, headers=headers, timeout=15)
            response.raise_for_status()  # Raise exception for bad status codes
            
            soup = BeautifulSoup(response.content, 'html.parser')
            
            # Remove script and style elements
            for script_or_style in soup(['script', 'style']):
                script_or_style.decompose()
            
            # Get text, preserving some structure
            text = soup.get_text(separator='\n', strip=True)
            
            # Try to get a good title
            title = soup.title.string if soup.title else url.split('/')[-1]
            
            return {
                "content": text,
                "filename": title,
                "extension": "url",
                "size": len(response.content),
                "char_count": len(text),
                "source_url": url
            }
        except requests.RequestException as e:
            raise Exception(f"Failed to fetch URL {url}: {e}")
        except Exception as e:
            raise Exception(f"Failed to parse URL {url}: {e}")
